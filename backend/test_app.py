import io
import os

os.environ.setdefault("GROQ_API_KEY", "test-key")
os.environ.setdefault("SECRET_KEY", "test-secret")

import pytest
from docx import Document
from fastapi.testclient import TestClient
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

import ai
import models  # noqa: F401  (registers the tables)
import readers
from database import Base, get_db
from main import app

engine = create_engine(
    "sqlite://", connect_args={"check_same_thread": False}, poolclass=StaticPool
)
TestingSession = sessionmaker(bind=engine, autoflush=False, autocommit=False)


def override_get_db():
    db = TestingSession()
    try:
        yield db
    finally:
        db.close()


app.dependency_overrides[get_db] = override_get_db

FAKE_QUIZ = [
    {"question": "Q1?", "options": ["a", "b", "c", "d"], "answer_index": 0, "explanation": "e1"},
    {"question": "Q2?", "options": ["a", "b", "c", "d"], "answer_index": 1, "explanation": "e2"},
]
NOTES = "A deadlock is a state where processes wait for each other forever."


@pytest.fixture(autouse=True)
def fresh_db():
    Base.metadata.drop_all(bind=engine)
    Base.metadata.create_all(bind=engine)
    yield


@pytest.fixture
def client():
    return TestClient(app)


def signup(client, email="ali@test.com", password="123456"):
    res = client.post("/auth/signup", json={"name": "Ali", "email": email, "password": password})
    assert res.status_code == 200
    return {"Authorization": f"Bearer {res.json()['access_token']}"}


def paste_doc(client, headers, title="OS", text=NOTES):
    res = client.post("/documents/paste", json={"title": title, "text": text}, headers=headers)
    assert res.status_code == 200
    return res.json()["id"]


def make_docx_bytes():
    doc = Document()
    doc.add_paragraph("Operating systems manage hardware and software resources.")
    doc.add_paragraph("A process is a program in execution.")
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


# ---------- Auth ----------
def test_signup_and_login(client):
    signup(client)
    res = client.post("/auth/login", data={"username": "ali@test.com", "password": "123456"})
    assert res.status_code == 200
    assert "access_token" in res.json()


def test_duplicate_email_rejected(client):
    signup(client)
    res = client.post("/auth/signup", json={"name": "B", "email": "ali@test.com", "password": "123456"})
    assert res.status_code == 400


def test_short_password_rejected(client):
    res = client.post("/auth/signup", json={"name": "A", "email": "a@test.com", "password": "123"})
    assert res.status_code == 400


def test_wrong_password_rejected(client):
    signup(client)
    res = client.post("/auth/login", data={"username": "ali@test.com", "password": "wrong"})
    assert res.status_code == 401


def test_protected_route_needs_token(client):
    assert client.get("/documents").status_code == 401


# ---------- Documents ----------
def test_paste_and_list_documents(client):
    headers = signup(client)
    doc_id = paste_doc(client, headers)
    res = client.get("/documents", headers=headers)
    assert res.status_code == 200
    assert [d["id"] for d in res.json()] == [doc_id]


def test_user_cannot_see_other_users_document(client):
    owner = signup(client, "owner@test.com")
    other = signup(client, "other@test.com")
    doc_id = paste_doc(client, owner)
    assert client.get(f"/documents/{doc_id}", headers=other).status_code == 404


def test_docx_upload_and_pages(client):
    headers = signup(client)
    files = {
        "file": (
            "notes.docx",
            make_docx_bytes(),
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )
    }
    res = client.post("/documents/upload", files=files, headers=headers)
    assert res.status_code == 200
    doc_id = res.json()["id"]
    pages = client.get(f"/documents/{doc_id}/pages", headers=headers).json()
    assert "process" in pages[0]["text"]


def test_unsupported_file_rejected(client):
    headers = signup(client)
    files = {"file": ("virus.exe", b"data", "application/octet-stream")}
    res = client.post("/documents/upload", files=files, headers=headers)
    assert res.status_code == 400


# ---------- Chat ----------
def test_chat_returns_answer_and_sources(client, monkeypatch):
    monkeypatch.setattr(ai, "ask_llm", lambda *a, **k: "Processes wait forever.\nPAGES: 1")
    headers = signup(client)
    doc_id = paste_doc(client, headers)
    res = client.post(
        f"/documents/{doc_id}/chat",
        json={"question": "What is a deadlock?", "style": "exam"},
        headers=headers,
    )
    assert res.status_code == 200
    assert res.json()["answer"] == "Processes wait forever."
    assert res.json()["sources"] == [1]
    history = client.get(f"/documents/{doc_id}/chat", headers=headers).json()
    assert len(history) == 2


# ---------- Quiz ----------
def test_quiz_scoring(client, monkeypatch):
    monkeypatch.setattr(ai, "generate_quiz", lambda *a, **k: FAKE_QUIZ)
    headers = signup(client)
    doc_id = paste_doc(client, headers)
    quiz = client.post(f"/documents/{doc_id}/quiz?count=2", headers=headers).json()
    assert "answer_index" not in quiz["questions"][0]  # answers stay hidden
    res = client.post(f"/quizzes/{quiz['quiz_id']}/submit", json={"answers": [0, 0]}, headers=headers)
    body = res.json()
    assert body["score"] == 1 and body["total"] == 2
    assert body["results"][0]["correct"] is True
    assert body["results"][1]["correct"] is False


# ---------- Flashcards ----------
def test_flashcard_spaced_repetition(client, monkeypatch):
    monkeypatch.setattr(
        ai, "generate_flashcards", lambda *a, **k: [{"question": "Q", "answer": "A"}] * 2
    )
    headers = signup(client)
    doc_id = paste_doc(client, headers)
    cards = client.post(f"/documents/{doc_id}/flashcards?count=2", headers=headers).json()
    card_id = cards[0]["id"]
    good = client.post(f"/flashcards/{card_id}/review", json={"correct": True}, headers=headers)
    assert good.json()["box"] == 2
    bad = client.post(f"/flashcards/{card_id}/review", json={"correct": False}, headers=headers)
    assert bad.json()["box"] == 1


# ---------- Dashboard ----------
def test_dashboard_counts(client):
    headers = signup(client)
    paste_doc(client, headers)
    data = client.get("/dashboard", headers=headers).json()
    assert data["documents"] == 1
    assert data["quizzes_taken"] == 0


# ---------- Helpers ----------
def test_split_answer_reads_pages_line():
    answer, pages = ai.split_answer("Hello there\nPAGES: 2, 4")
    assert answer == "Hello there"
    assert pages == [2, 4]


def test_split_answer_without_pages_line():
    answer, pages = ai.split_answer("Just an answer")
    assert answer == "Just an answer"
    assert pages is None


def test_pagination_splits_long_text():
    text = "\n\n".join(["word " * 100] * 10)
    pages = readers.text_to_pages(text)
    assert len(pages) >= 2
    assert [p for p, _ in pages] == list(range(1, len(pages) + 1))
    assert all(len(t) <= 2800 for _, t in pages)